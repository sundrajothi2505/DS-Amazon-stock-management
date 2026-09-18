import csv
import io
import re
from datetime import datetime, date

import openpyxl

# Column names we recognize, matched case-insensitively against whatever
# header row the uploaded file actually has. This lets the same importer
# handle Amazon's "All Listings Report", "Active Listings Report", or a
# plain custom CSV/TSV with differently-ordered columns.
SKU_COLUMNS = {"seller-sku", "sku", "seller sku"}
NAME_COLUMNS = {"item-name", "name", "product-name", "item name", "product name"}
ASIN_COLUMNS = {"asin1", "asin"}
QUANTITY_COLUMNS = {"quantity", "stock", "current-stock", "qty"}
STATUS_COLUMNS = {"status"}


def _find_column(header, candidates):
    for i, col in enumerate(header):
        if col in candidates:
            return i
    return None


def parse_report(raw_bytes):
    """Parse an Amazon seller listings report (tab- or comma-separated).

    Returns (rows, error). On success, rows is a list of dicts with keys
    sku, name, asin, quantity (int or None if the file had no quantity
    column or it wasn't a number), and status. On failure, rows is [] and
    error explains what went wrong.
    """
    try:
        text = raw_bytes.decode("utf-8-sig", errors="replace")
    except Exception:
        return [], "Couldn't read that file — is it a text/CSV export?"

    lines = text.splitlines()
    if not lines:
        return [], "The file appears to be empty."

    delimiter = "\t" if "\t" in lines[0] else ","
    reader = csv.reader(io.StringIO(text), delimiter=delimiter, quoting=csv.QUOTE_NONE)

    try:
        header_row = next(reader)
    except StopIteration:
        return [], "The file appears to be empty."
    header = [h.strip().lower() for h in header_row]

    sku_idx = _find_column(header, SKU_COLUMNS)
    name_idx = _find_column(header, NAME_COLUMNS)
    asin_idx = _find_column(header, ASIN_COLUMNS)
    qty_idx = _find_column(header, QUANTITY_COLUMNS)
    status_idx = _find_column(header, STATUS_COLUMNS)

    if sku_idx is None or name_idx is None:
        return [], (
            "Couldn't find SKU and product name columns in this file. "
            "This importer expects an Amazon listings report (with columns "
            "like 'seller-sku' and 'item-name') — is that what you uploaded?"
        )

    def get(raw, idx):
        return raw[idx].strip() if idx is not None and idx < len(raw) else ""

    rows = []
    for raw in reader:
        if not raw or all(not c.strip() for c in raw):
            continue
        sku = get(raw, sku_idx)
        if not sku:
            continue
        quantity = None
        if qty_idx is not None:
            raw_qty = get(raw, qty_idx)
            try:
                quantity = int(float(raw_qty))
            except ValueError:
                quantity = None
        rows.append({
            "sku": sku,
            "name": get(raw, name_idx),
            "asin": get(raw, asin_idx) if asin_idx is not None else "",
            "quantity": quantity,
            "status": get(raw, status_idx) if status_idx is not None else "",
        })

    return rows, None


# ---------------------------------------------------------------------------
# Order history import from a messy multi-sheet Excel workbook.
#
# Real seller spreadsheets like this rarely have a clean single header row:
# there's usually a title above it, the header row lands on a different row
# number in every sheet, column names are spelled slightly differently month
# to month ("ORDER DATE" vs "ORDER -DATE"), and the same order can appear in
# more than one sheet (a monthly tab plus a running "returns" or scratch
# tab). This scans every sheet for anything that looks like an order table,
# rather than assuming a fixed layout.
# ---------------------------------------------------------------------------

def _normalize_header(value):
    return re.sub(r"[\s\-_]+", " ", str(value).strip().lower()).strip()


def _find_order_columns(row_values):
    """Given one row's cell values, return a column map if this row looks
    like an order-table header, else None."""
    col_map = {}
    has_order_id = has_order_date = has_desc = False
    for idx, value in enumerate(row_values):
        if value is None or str(value).strip() == "":
            continue
        norm = _normalize_header(value)
        if "order date" in norm:
            col_map["order_date"] = idx
            has_order_date = True
        elif "order id" in norm:
            col_map["order_id"] = idx
            has_order_id = True
        elif "discription" in norm or "description" in norm:
            col_map["description"] = idx
            has_desc = True
        elif "shipping date" in norm:
            col_map["shipping_date"] = idx
        elif "delivery status" in norm:
            col_map["delivery_status"] = idx
        elif "return" in norm and "cancel" in norm:
            col_map["return_cancel"] = idx
        elif "qty" in norm:
            col_map["qty"] = idx
    if has_order_id and (has_order_date or has_desc):
        return col_map
    return None


def _cell_to_date_str(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    text = str(value).strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _cell_to_order_number(value):
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell_to_qty(value):
    try:
        q = int(float(value))
        return q if q > 0 else 1
    except (TypeError, ValueError):
        return 1


def map_order_status(delivery_status_text, return_cancel_text):
    """Collapse a seller's own free-text status columns down to this app's
    fixed status set, without losing the original wording (callers should
    keep it in the row's notes)."""
    d = (delivery_status_text or "").strip().lower()
    r = (return_cancel_text or "").strip().lower()
    if "return" in r:
        return "returned"
    if "cancel" in r:
        return "cancelled"
    if "lost" in d or "lost" in r:
        return "lost"
    if "damage" in d or "damage" in r:
        return "damaged"
    if "deliver" in d:
        return "delivered"
    return "shipped"


def parse_orders_workbook(raw_bytes, max_header_search_rows=15):
    """Parse every sheet of an uploaded .xlsx for order-shaped tables.

    Returns (rows, sheet_report). rows is a list of dicts: order_number,
    order_date, description, quantity, shipped_date, status, status_note.
    sheet_report is a list of (sheet_name, message) describing what was
    found in each sheet, for a transparent import summary — including
    sheets that were skipped because they didn't look like an order table.
    Rows are de-duplicated by order_number *within this file*, preferring
    whichever sheet is processed first (monthly-named sheets are processed
    before generic/scratch sheets like "Sheet1" or a consolidated "return"
    tab, since those tend to be redundant copies).
    """
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw_bytes), data_only=True, read_only=True)
    except Exception:
        return [], [("(file)", "Couldn't open this as an Excel file.")]

    def sheet_priority(name):
        low = name.lower()
        if "return" in low or low.strip() == "sheet1":
            return 1
        return 0

    sheet_names = sorted(wb.sheetnames, key=sheet_priority)

    rows = []
    seen_order_numbers = set()
    sheet_report = []

    for sheet_name in sheet_names:
        ws = wb[sheet_name]
        col_map = None
        header_row_num = None
        for i, row in enumerate(ws.iter_rows(min_row=1, max_row=max_header_search_rows, values_only=True)):
            candidate = _find_order_columns(row)
            if candidate:
                col_map = candidate
                header_row_num = i + 1
                break

        if col_map is None:
            sheet_report.append((sheet_name, "No order table found — skipped."))
            continue

        added_from_sheet = 0
        skipped_from_sheet = 0
        dup_from_sheet = 0

        for row in ws.iter_rows(min_row=header_row_num + 1, values_only=True):
            def get(key):
                idx = col_map.get(key)
                return row[idx] if idx is not None and idx < len(row) else None

            order_number = _cell_to_order_number(get("order_id"))
            if not order_number:
                continue

            order_date = _cell_to_date_str(get("order_date"))
            if not order_date:
                skipped_from_sheet += 1
                continue

            if order_number in seen_order_numbers:
                dup_from_sheet += 1
                continue
            seen_order_numbers.add(order_number)

            delivery_status_raw = get("delivery_status")
            return_cancel_raw = get("return_cancel")
            status = map_order_status(
                str(delivery_status_raw) if delivery_status_raw else "",
                str(return_cancel_raw) if return_cancel_raw else "",
            )
            status_bits = [str(v).strip() for v in (delivery_status_raw, return_cancel_raw) if v and str(v).strip()]

            rows.append({
                "order_number": order_number,
                "order_date": order_date,
                "description": (str(get("description")).strip() if get("description") else "") or None,
                "quantity": _cell_to_qty(get("qty")),
                "shipped_date": _cell_to_date_str(get("shipping_date")),
                "status": status,
                "status_note": " / ".join(status_bits) if status_bits else None,
            })
            added_from_sheet += 1

        msg = f"{added_from_sheet} order(s) found"
        extras = []
        if dup_from_sheet:
            extras.append(f"{dup_from_sheet} duplicate(s) already seen elsewhere in this file")
        if skipped_from_sheet:
            extras.append(f"{skipped_from_sheet} row(s) skipped (no usable date)")
        if extras:
            msg += " — " + ", ".join(extras)
        sheet_report.append((sheet_name, msg))

    return rows, sheet_report
